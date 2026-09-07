"""The editable Word (.docx) companion of the authentic form."""

from __future__ import annotations

import docx

from barbarika_compliance import generate


def _fill(seeded_vault, submission_toml, tmp_path, monkeypatch):
    monkeypatch.setenv("COMPLIANCE_AUDIT_LOG", str(tmp_path / "audit.log.jsonl"))
    out_docx = str(tmp_path / "form.docx")
    return generate.generate_form_docx(
        db_path=seeded_vault["db"], submission_path=submission_toml,
        out_docx=out_docx, incident="latest", key_path=seeded_vault["key"],
    )


def test_form_docx_fills_authentic_cells(seeded_vault, submission_toml, tmp_path, monkeypatch):
    res = _fill(seeded_vault, submission_toml, tmp_path, monkeypatch)
    d = docx.Document(res["docx"])
    cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    # Authentic form labels are still present (we filled the real template)...
    assert "Contact Information of the Reporter" in cells
    assert "Name & Role/Title" in cells
    assert "Incident Type" in cells
    # ...and our values landed in the cells.
    assert seeded_vault["reviewer"] in cells             # Name & Role/Title
    assert seeded_vault["attacker_ip"] in cells          # IP Address
    assert "Meridian FinServ" in cells                   # Organization name
    assert "ciso@meridianfinserv.in" in cells            # Email


def test_form_docx_is_a_real_docx(seeded_vault, submission_toml, tmp_path, monkeypatch):
    res = _fill(seeded_vault, submission_toml, tmp_path, monkeypatch)
    # Opens as a valid single-table Word document.
    d = docx.Document(res["docx"])
    assert len(d.tables) == 1
    assert len(d.tables[0].rows) == 15
